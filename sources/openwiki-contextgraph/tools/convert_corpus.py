"""
Convert the raw mock documents (PDF, XLSX, DOCX) into Markdown that OpenWiki's agent can read.

OpenWiki's documentation agent reads text files with shell tools; it cannot open binary
PDF/Excel/Word files. This script writes a text corpus into <out>/sources/ and keeps the
originals in <out>/raw/ (excluded from OpenWiki via .openwikiignore).

Usage:
    python tools/convert_corpus.py --data data --out semantic-corpus

Requires: pymupdf4llm, openpyxl, markitdown  (pip install pymupdf4llm openpyxl "markitdown[docx]")
Cross-platform (Windows / macOS / Linux).
"""
import argparse
import re
import shutil
from pathlib import Path

import pymupdf4llm
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


# ------------------------------------------------------------------ PDF
_NOISE = re.compile(r"^\s*(\**MERIDIAN HARBOR FINANCIAL CORP\.\**|Synthetic document - fictional institution.*|\d{1,3})\s*$")


def _clean_pdf_page(text: str) -> str:
    """Drop running headers/footers, page numbers and dot-leader table-of-contents rows."""
    out = []
    lines = text.splitlines()
    for idx, line in enumerate(lines):
        if _NOISE.match(line):
            continue
        if line.count(". ") > 8:          # table-of-contents dot leaders
            continue
        if idx < 6 and len(line.strip()) < 70 and not line.startswith(("#", "|")) and \
                re.search(r"(Annual Report|Earnings Release|Pillar 3 Regulatory Capital Disclosures -)", line):
            continue                      # running header on the right
        out.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out))


def convert_pdf(pdf: Path, out_dir: Path) -> Path:
    pages = pymupdf4llm.to_markdown(str(pdf), page_chunks=True)
    parts = [f"<!-- source: raw/reports/{pdf.name} | converted by tools/convert_corpus.py -->",
             f"# {pdf.stem.replace('_', ' ')}", ""]
    for i, page in enumerate(pages, 1):
        text = page["text"] if isinstance(page, dict) else str(page)
        parts.append(f"\n<!-- page {i} -->\n")
        parts.append(_clean_pdf_page(text).strip())
    target = out_dir / f"{slug(pdf.stem)}.md"
    target.write_text("\n".join(parts), encoding="utf-8")
    return target


# ------------------------------------------------------------------ XLSX
def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if abs(v) < 1 and v != 0:
            return f"{v:.4f}"
        return f"{v:,.2f}"
    if isinstance(v, int):
        return f"{v:,}"
    return str(v).replace("|", "/").replace("\n", " ")


def convert_xlsx(xlsx: Path, out_dir: Path) -> Path:
    wb_f = load_workbook(xlsx)                   # formulas
    wb_v = load_workbook(xlsx, data_only=True)   # cached values (file must be recalculated in Excel/LibreOffice)
    lines = [f"<!-- source: raw/models/{xlsx.name} | converted by tools/convert_corpus.py -->",
             f"# Excel model: {xlsx.name}", "",
             "This file is a text rendering of an Excel workbook. Each sheet is shown as a value grid, followed "
             "by the list of formulas in that sheet (cell, row label, column header, formula, computed value). "
             "Blue-font cells in the workbook are inputs; black cells are formulas; green cells link to other sheets.",
             ""]
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        lines += [f"## Sheet: {ws_f.title}", ""]
        max_r, max_c = ws_f.max_row, ws_f.max_column
        # README-style sheets: render as key/value text
        if ws_f.title.upper() == "README":
            for r in range(1, max_r + 1):
                a, b = ws_v.cell(r, 1).value, ws_v.cell(r, 2).value
                if a and b:
                    lines.append(f"- **{_fmt(a)}** {_fmt(b)}")
                elif a:
                    lines.append(f"\n### {_fmt(a)}\n" if len(str(a)) < 60 else _fmt(a))
                elif b:
                    lines.append(f"- {_fmt(b)}")
            lines.append("")
            continue
        # title rows (1-3) as text
        for r in range(1, 4):
            v = ws_v.cell(r, 1).value
            if v:
                lines.append(f"> {_fmt(v)}")
        lines.append("")
        # value grid from row 5
        header_row = 5
        cols = [c for c in range(1, max_c + 1) if any(ws_v.cell(r, c).value is not None for r in range(header_row, max_r + 1))]
        if cols:
            lines.append("| Row | " + " | ".join(f"{get_column_letter(c)}" for c in cols) + " |")
            lines.append("|---|" + "---|" * len(cols))
            for r in range(header_row, max_r + 1):
                vals = [_fmt(ws_v.cell(r, c).value) for c in cols]
                if any(vals):
                    lines.append(f"| {r} | " + " | ".join(vals) + " |")
            lines.append("")
        # formulas
        flist = []
        for row in ws_f.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    label = ws_v.cell(c.row, 1).value
                    hdr = ws_v.cell(header_row, c.column).value
                    flist.append((c.coordinate, _fmt(label), _fmt(hdr), c.value.replace("|", "/"),
                                  _fmt(ws_v[c.coordinate].value)))
        if flist:
            lines += [f"### Formulas in {ws_f.title} ({len(flist)})", "",
                      "| Cell | Row label | Column header | Formula | Value |", "|---|---|---|---|---|"]
            lines += [f"| {a} | {b} | {c} | `{d}` | {e} |" for a, b, c, d, e in flist]
            lines.append("")
        # comments (assumption sources)
        notes = [(c.coordinate, c.comment.text) for row in ws_f.iter_rows() for c in row if c.comment]
        if notes:
            lines += [f"### Cell notes in {ws_f.title}", ""] + [f"- {a}: {t}" for a, t in notes] + [""]
    target = out_dir / f"{slug(xlsx.stem)}.md"
    target.write_text("\n".join(lines), encoding="utf-8")
    return target


# ------------------------------------------------------------------ DOCX
def convert_docx(docx: Path, out_dir: Path) -> list:
    from markitdown import MarkItDown
    md = MarkItDown().convert(str(docx)).text_content
    header = f"<!-- source: raw/api_docs/{docx.name} | converted by tools/convert_corpus.py -->\n"
    full = out_dir / f"{slug(docx.stem)}.md"
    full.write_text(header + md, encoding="utf-8")
    written = [full]
    # also split per API so each API gets its own source file
    per_api = out_dir / "apis"
    per_api.mkdir(exist_ok=True)
    chunks = re.split(r"(?m)^(?=#{1,3} API-\d\d )", md)
    for ch in chunks:
        m = re.match(r"#{1,3} (API-\d\d) (.+)", ch)
        if not m:
            continue
        body = re.split(r"(?m)^#{1,3} Appendix", ch)[0]
        p = per_api / f"{m.group(1).lower()}-{slug(m.group(2))}.md"
        p.write_text(header + body.strip() + "\n", encoding="utf-8")
        written.append(p)
    return written


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default="data", help="folder with reports/, models/, api_docs/")
    ap.add_argument("--out", default="semantic-corpus", help="corpus repository root")
    args = ap.parse_args()
    data, out = Path(args.data), Path(args.out)
    src = out / "sources"
    for sub in ("reports", "models", "api_docs"):
        (src / sub).mkdir(parents=True, exist_ok=True)
        (out / "raw" / sub).mkdir(parents=True, exist_ok=True)
    n = 0
    for pdf in sorted((data / "reports").glob("*.pdf")):
        shutil.copy2(pdf, out / "raw" / "reports" / pdf.name)
        print("pdf  ->", convert_pdf(pdf, src / "reports")); n += 1
    for x in sorted((data / "models").glob("*.xlsx")):
        shutil.copy2(x, out / "raw" / "models" / x.name)
        print("xlsx ->", convert_xlsx(x, src / "models")); n += 1
    for d in sorted((data / "api_docs").glob("*.docx")):
        shutil.copy2(d, out / "raw" / "api_docs" / d.name)
        for p in convert_docx(d, src / "api_docs"):
            print("docx ->", p)
        n += 1
    print(f"converted {n} source documents into {src}")


if __name__ == "__main__":
    main()
