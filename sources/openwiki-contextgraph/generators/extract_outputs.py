"""Read recalculated model outputs so the PDFs quote the same numbers as the Excel models."""
import json
from openpyxl import load_workbook

D = "../data/models/"
out = {}

w = load_workbook(D + "NII_Sensitivity_Model.xlsx", data_only=True)
s = w["Summary"]
out["nii"] = {s[f"A{r}"].value: {"shock": s[f"B{r}"].value, "nii": s[f"C{r}"].value, "delta": s[f"D{r}"].value,
                                 "pct": s[f"E{r}"].value, "eve": s[f"G{r}"].value, "eve_pct": s[f"H{r}"].value}
              for r in range(6, 11)}

w = load_workbook(D + "CECL_Allowance_Model.xlsx", data_only=True)
a = w["Allowance_Summary"]
out["cecl"] = {a[f"A{r}"].value: {"upside": a[f"B{r}"].value, "baseline": a[f"C{r}"].value,
                                  "downside": a[f"D{r}"].value, "modeled": a[f"E{r}"].value,
                                  "overlay": a[f"F{r}"].value, "total": a[f"G{r}"].value,
                                  "coverage": a[f"K{r}"].value}
               for r in range(6, 13)}
se = w["Sensitivity"]
out["cecl_sens"] = {se[f"A{r}"].value: {"allowance": se[f"C{r}"].value, "delta": se[f"D{r}"].value,
                                        "pct": se[f"E{r}"].value} for r in range(6, 10)}

w = load_workbook(D + "Capital_Planning_Model.xlsx", data_only=True)
s = w["Summary"]
out["capital"] = {s[f"A{r}"].value: s[f"B{r}"].value for r in range(6, 19)}
for name in ("Baseline_Projection", "Stress_Projection"):
    ws = w[name]
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row, values_only=True):
        if row[0] == "CET1 ratio":
            out[name + "_cet1"] = list(row[1:11])

json.dump(out, open("model_outputs.json", "w"), indent=1, default=str)
print(json.dumps(out, indent=1, default=str)[:3000])
