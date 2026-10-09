"""Build the three Excel financial models with live formulas."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
import bankdata as B

OUT = sys.argv[1] if len(sys.argv) > 1 else "../data/models"

ARIAL = "Arial"
BLUE = Font(name=ARIAL, color="0000FF", size=10)
BLACK = Font(name=ARIAL, color="000000", size=10)
GREEN = Font(name=ARIAL, color="008000", size=10)
BOLD = Font(name=ARIAL, bold=True, size=10)
HDR = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
TITLE = Font(name=ARIAL, bold=True, size=14, color="0B2545")
SUB = Font(name=ARIAL, italic=True, size=9, color="555555")
NAVY = PatternFill("solid", fgColor="0B2545")
YELLOW = PatternFill("solid", fgColor="FFFF00")
LIGHT = PatternFill("solid", fgColor="EEF2F7")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
TOPLINE = Border(top=Side(style="thin", color="000000"))

USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
BPS = '0" bps";(0" bps");"0 bps"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '0.00;(0.00);"-"'


def title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = SUB
    ws["A3"] = B.BANK["disclaimer"]
    ws["A3"].font = Font(name=ARIAL, size=8, color="AA0000", italic=True)


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = HDR
        c.fill = NAVY
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, note=None, bold=False):
    c = ws[ref]
    c.value = value
    f = Font(name=font.name, color=font.color, size=font.size, bold=bold or font.bold)
    c.font = f
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, "Model owner")
    c.border = BOX
    return c


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def readme(wb, model, sheets, how_to, inputs, outputs, limits):
    ws = wb.active
    ws.title = "README"
    title(ws, f"{model['name']}  ({model['id']})", f"{B.BANK['name']} - model documentation sheet")
    rows = [
        ("Model ID", model["id"]), ("Model name", model["name"]), ("Risk tier", model["tier"]),
        ("Model owner", model["owner"]), ("Independent validator", model["validator"]),
        ("Last validation", model["last_validation"]), ("Next validation due", model["next_validation"]),
        ("Purpose", model["purpose"]),
        ("Upstream data feeds (APIs)", "; ".join(model["apis"])),
        ("Downstream disclosures", "; ".join(model["reports"])),
        ("Units", "USD millions unless stated otherwise; rates stored as decimals"),
        ("As-of date", "See Assumptions sheet"),
    ]
    r = 5
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = BOLD
        c = ws.cell(row=r, column=2, value=v)
        c.font = BLACK
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Sheets").font = TITLE
    r += 1
    for name, desc in sheets:
        ws.cell(row=r, column=1, value=name).font = BOLD
        ws.cell(row=r, column=2, value=desc).alignment = Alignment(wrap_text=True)
        r += 1
    for heading, items in (("How the model works", how_to), ("Key inputs", inputs),
                           ("Key outputs", outputs), ("Limits, controls and known limitations", limits)):
        r += 1
        ws.cell(row=r, column=1, value=heading).font = TITLE
        r += 1
        for i, item in enumerate(items, 1):
            ws.cell(row=r, column=1, value=f"{i}.").font = BLACK
            c = ws.cell(row=r, column=2, value=item)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    r += 1
    ws.cell(row=r, column=1, value="Colour legend").font = TITLE
    r += 1
    for txt, font, fill in (("Blue text = hard-coded input / scenario lever", BLUE, None),
                            ("Black text = formula", BLACK, None),
                            ("Green text = link to another sheet", GREEN, None),
                            ("Yellow fill = key assumption", BLACK, YELLOW)):
        c = ws.cell(row=r, column=2, value=txt)
        c.font = font
        if fill:
            c.fill = fill
        r += 1
    widths(ws, {"A": 28, "B": 110})
    for row in ws.iter_rows(min_row=5, max_row=r):
        for c in row:
            if c.column == 2:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


# =============================================================== NII MODEL
def build_nii():
    m = B.MODELS[0]
    wb = Workbook()
    readme(wb, m,
           [("Assumptions", "As-of date, policy rate, rate shock grid, repricing timing conventions, risk limits."),
            ("Balance_Sheet", "Rate-sensitive assets and liabilities with balances, yields/costs, repricing profile, "
                              "deposit betas and modified durations."),
            ("NII_Projection", "Base-case 12-month NII and the change in NII for each parallel shock."),
            ("EVE", "Economic value of equity sensitivity using a modified-duration approximation."),
            ("Summary", "Results versus Board-approved IRRBB limits; feeds Annual Report and Pillar 3.")],
           ["Base NII = sum over assets of balance x yield, minus sum over liabilities of balance x cost.",
            "For each shock, every position reprices only for the fraction of the 12-month horizon remaining "
            "after its repricing date (bucket midpoint): 0-3m reprices at month 1.5, 3-12m at month 7.5.",
            "Asset rates move one-for-one with the shock (asset beta in column L); liability rates move by "
            "shock x deposit beta.",
            "Change in NII = asset repricing effect - liability repricing effect. Static balance sheet "
            "(no growth, no mix shift).",
            "EVE change = -(asset value x modified duration - liability value x modified duration) x shock."],
           ["Balances and yields from API-15 Treasury Liquidity Positions API (month-end snapshot).",
            "Loan repricing profiles from API-11 Loan Servicing API.",
            "Yield curve and policy rate from API-12 Market Data API.",
            "Deposit betas from the Deposit Behaviour sub-model (annual calibration, MRGR approved)."],
           ["12-month NII under +/-100 bp and +/-200 bp parallel shocks.",
            "EVE sensitivity as a % of Tier 1 capital.",
            "Limit utilisation flags used in ALCO reporting."],
           ["Board limit: NII decline no greater than 7.0% of base NII in any parallel shock.",
            "Board limit: EVE decline no greater than 15.0% of Tier 1 capital.",
            "Limitations: parallel shocks only (no twists), static balance sheet, no rate floors on "
            "deposit costs below zero, betas held constant across shock sizes.",
            "Compensating control: dynamic balance-sheet simulation run quarterly in the ALM engine."])

    # ---------------- Assumptions
    a = wb.create_sheet("Assumptions")
    title(a, "Assumptions", "All rates as decimals; shocks in basis points")
    rows = [
        ("As-of date", "2025-12-31", None, "Month-end snapshot used for the FY2025 Annual Report."),
        ("Policy rate (fed funds upper bound)", 0.0375, PCT2, "Source: API-12 Market Data API, series POLICY.FEDFUNDS.UB"),
        ("Projection horizon (months)", 12, '0', "Standard IRRBB earnings horizon."),
        ("Repricing month - bucket 0-3m (midpoint)", 1.5, NUM1, "Convention: midpoint of bucket."),
        ("Repricing month - bucket 3-12m (midpoint)", 7.5, NUM1, "Convention: midpoint of bucket."),
        ("NII limit (max decline, % of base)", 0.07, PCT, "Board Risk Committee limit, approved 2025-03."),
        ("EVE limit (max decline, % of Tier 1)", 0.15, PCT, "Board Risk Committee limit, approved 2025-03."),
        ("Tier 1 capital ($mm)", B.CAPITAL["Tier 1 capital"][0], USD, "Source: API-16 Regulatory Reporting API, YE2025."),
    ]
    header(a, 5, ["Assumption", "Value", "Source / note"])
    for i, (k, v, fmt, note) in enumerate(rows):
        r = 6 + i
        put(a, f"A{r}", k, BOLD)
        put(a, f"B{r}", v, BLUE, fmt, YELLOW if i in (1, 5, 6) else None)
        put(a, f"C{r}", note, SUB)
    # shock grid
    header(a, 16, ["Scenario", "Shock (bps)"])
    shocks = [("Down 200", -200), ("Down 100", -100), ("Base", 0), ("Up 100", 100), ("Up 200", 200)]
    for i, (n, s) in enumerate(shocks):
        put(a, f"A{17+i}", n, BOLD)
        put(a, f"B{17+i}", s, BLUE, BPS)
    widths(a, {"A": 42, "B": 16, "C": 70})

    # ---------------- Balance sheet
    bs = wb.create_sheet("Balance_Sheet")
    title(bs, "Rate-Sensitive Balance Sheet", "Balances $mm at 2025-12-31; repricing shares sum to 100% per row")
    cols = ["Line", "Side", "Balance ($mm)", "Yield / cost", "Reprice 0-3m", "Reprice 3-12m",
            "Reprice 1-5y", "Reprice >5y", "Share check", "Annual interest ($mm)", "Mod. duration (yrs)",
            "Rate beta", "Feed"]
    header(bs, 5, cols)
    lines = [
        # name, side, balance, yield, b03, b312, b15, b5p, duration, beta, feed
        ("Deposits with banks & fed funds sold", "Asset", 246_900, 0.0372, 1.00, 0.00, 0.00, 0.00, 0.1, 1.00, "API-15"),
        ("Investment securities (AFS + HTM)", "Asset", 170_900, 0.0386, 0.12, 0.10, 0.38, 0.40, 4.6, 1.00, "API-15"),
        ("Trading assets (interest-earning)", "Asset", 98_400, 0.0410, 0.35, 0.25, 0.25, 0.15, 2.1, 1.00, "API-15"),
        ("Credit card loans", "Asset", 138_400, 0.1650, 0.62, 0.08, 0.30, 0.00, 0.9, 1.00, "API-11"),
        ("Residential mortgage loans", "Asset", 218_600, 0.0428, 0.05, 0.07, 0.28, 0.60, 5.8, 1.00, "API-11"),
        ("Auto loans", "Asset", 64_900, 0.0655, 0.06, 0.24, 0.70, 0.00, 1.7, 1.00, "API-11"),
        ("Commercial real estate loans", "Asset", 98_200, 0.0640, 0.55, 0.10, 0.30, 0.05, 1.6, 1.00, "API-11"),
        ("Commercial & industrial loans", "Asset", 172_500, 0.0640, 0.78, 0.08, 0.12, 0.02, 0.6, 1.00, "API-11"),
        ("Other consumer & wholesale loans", "Asset", 49_700, 0.0590, 0.50, 0.15, 0.30, 0.05, 1.4, 1.00, "API-11"),
        ("Consumer interest-bearing deposits", "Liability", 402_000, 0.0205, 0.70, 0.30, 0.00, 0.00, 2.8, 0.45, "API-15"),
        ("Wholesale interest-bearing deposits", "Liability", 338_400, 0.0315, 0.85, 0.15, 0.00, 0.00, 0.6, 0.75, "API-15"),
        ("Noninterest-bearing deposits", "Liability", 272_000, 0.0000, 0.00, 0.00, 0.40, 0.60, 3.5, 0.00, "API-15"),
        ("Repo & short-term borrowings", "Liability", 81_000, 0.0395, 1.00, 0.00, 0.00, 0.00, 0.1, 1.00, "API-15"),
        ("Long-term debt", "Liability", 86_700, 0.0482, 0.30, 0.10, 0.35, 0.25, 4.9, 1.00, "API-15"),
    ]
    first = 6
    for i, ln in enumerate(lines):
        r = first + i
        name, side, bal, y, b1, b2, b3, b4, dur, beta, feed = ln
        put(bs, f"A{r}", name, BOLD)
        put(bs, f"B{r}", side, BLUE)
        put(bs, f"C{r}", bal, BLUE, USD)
        put(bs, f"D{r}", y, BLUE, PCT2)
        for col, v in zip("EFGH", (b1, b2, b3, b4)):
            put(bs, f"{col}{r}", v, BLUE, PCT)
        put(bs, f"I{r}", f'=IF(ABS(SUM(E{r}:H{r})-1)<0.0001,"OK","CHECK")', BLACK)
        put(bs, f"J{r}", f"=C{r}*D{r}", BLACK, USD)
        put(bs, f"K{r}", dur, BLUE, NUM1)
        put(bs, f"L{r}", beta, BLUE, NUM2, YELLOW if side == "Liability" else None,
            note="Deposit beta from Deposit Behaviour sub-model (2025 calibration)." if side == "Liability" else None)
        put(bs, f"M{r}", feed, BLUE)
    last = first + len(lines) - 1
    tr = last + 2
    put(bs, f"A{tr}", "Total rate-sensitive assets", BOLD)
    put(bs, f"C{tr}", f'=SUMIFS(C{first}:C{last},B{first}:B{last},"Asset")', BLACK, USD, bold=True)
    put(bs, f"J{tr}", f'=SUMIFS(J{first}:J{last},B{first}:B{last},"Asset")', BLACK, USD, bold=True)
    put(bs, f"A{tr+1}", "Total rate-sensitive liabilities", BOLD)
    put(bs, f"C{tr+1}", f'=SUMIFS(C{first}:C{last},B{first}:B{last},"Liability")', BLACK, USD, bold=True)
    put(bs, f"J{tr+1}", f'=SUMIFS(J{first}:J{last},B{first}:B{last},"Liability")', BLACK, USD, bold=True)
    put(bs, f"A{tr+2}", "Base-case annual net interest income", BOLD)
    put(bs, f"J{tr+2}", f"=J{tr}-J{tr+1}", BLACK, USD, bold=True)
    widths(bs, {"A": 38, "B": 10, "C": 14, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 9,
                "J": 15, "K": 12, "L": 10, "M": 9})
    bs.freeze_panes = "B6"

    # ---------------- NII projection
    n = wb.create_sheet("NII_Projection")
    title(n, "12-Month NII Projection by Parallel Shock",
          "Effect = balance x shock x beta x (share 0-3m x (H-1.5)/H + share 3-12m x (H-7.5)/H)")
    header(n, 5, ["Line", "Side", "Repricing factor"] + [s[0] for s in shocks])
    put(n, "A6", "Shock (bps)", BOLD)
    for j in range(5):
        col = get_column_letter(4 + j)
        put(n, f"{col}6", f"=Assumptions!B{17+j}", GREEN, BPS)
    for i in range(len(lines)):
        r = 7 + i
        src = first + i
        put(n, f"A{r}", f"=Balance_Sheet!A{src}", GREEN)
        put(n, f"B{r}", f"=Balance_Sheet!B{src}", GREEN)
        put(n, f"C{r}",
            f"=Balance_Sheet!E{src}*(Assumptions!$B$8-Assumptions!$B$9)/Assumptions!$B$8"
            f"+Balance_Sheet!F{src}*(Assumptions!$B$8-Assumptions!$B$10)/Assumptions!$B$8",
            BLACK, '0.000')
        for j in range(5):
            col = get_column_letter(4 + j)
            sign = f'IF(B{r}="Asset",1,-1)'
            put(n, f"{col}{r}",
                f"={sign}*Balance_Sheet!C{src}*Balance_Sheet!L{src}*$C{r}*{col}$6/10000",
                BLACK, USD)
    nl = 7 + len(lines) - 1
    rr = nl + 2
    put(n, f"A{rr}", "Change in NII ($mm)", BOLD)
    put(n, f"A{rr+1}", "Base-case NII ($mm)", BOLD)
    put(n, f"A{rr+2}", "Projected NII ($mm)", BOLD)
    put(n, f"A{rr+3}", "Change in NII (%)", BOLD)
    for j in range(5):
        col = get_column_letter(4 + j)
        put(n, f"{col}{rr}", f"=SUM({col}7:{col}{nl})", BLACK, USD, bold=True)
        put(n, f"{col}{rr+1}", f"=Balance_Sheet!$J${tr+2}", GREEN, USD)
        put(n, f"{col}{rr+2}", f"={col}{rr+1}+{col}{rr}", BLACK, USD, bold=True)
        put(n, f"{col}{rr+3}", f"=IF({col}{rr+1}=0,0,{col}{rr}/{col}{rr+1})", BLACK, PCT)
    widths(n, {"A": 38, "B": 10, "C": 14, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13})
    nii_rows = {"delta": rr, "base": rr + 1, "proj": rr + 2, "pct": rr + 3}

    # ---------------- EVE
    e = wb.create_sheet("EVE")
    title(e, "Economic Value of Equity Sensitivity", "dEVE ~= -(sum asset PV x duration - sum liability PV x duration) x shock")
    header(e, 5, ["Line", "Side", "Balance ($mm)", "Mod. duration", "Dollar duration ($mm per 100 bp)"]
           + [s[0] for s in shocks])
    put(e, "A6", "Shock (bps)", BOLD)
    for j in range(5):
        col = get_column_letter(6 + j)
        put(e, f"{col}6", f"=Assumptions!B{17+j}", GREEN, BPS)
    for i in range(len(lines)):
        r = 7 + i
        src = first + i
        put(e, f"A{r}", f"=Balance_Sheet!A{src}", GREEN)
        put(e, f"B{r}", f"=Balance_Sheet!B{src}", GREEN)
        put(e, f"C{r}", f"=Balance_Sheet!C{src}", GREEN, USD)
        put(e, f"D{r}", f"=Balance_Sheet!K{src}", GREEN, NUM1)
        put(e, f"E{r}", f'=IF(B{r}="Asset",1,-1)*C{r}*D{r}/100', BLACK, USD)
        for j in range(5):
            col = get_column_letter(6 + j)
            put(e, f"{col}{r}", f"=-$E{r}*{col}$6/100", BLACK, USD)
    el = 7 + len(lines) - 1
    er = el + 2
    put(e, f"A{er}", "Change in EVE ($mm)", BOLD)
    put(e, f"A{er+1}", "Tier 1 capital ($mm)", BOLD)
    put(e, f"A{er+2}", "Change in EVE (% of Tier 1)", BOLD)
    for j in range(5):
        col = get_column_letter(6 + j)
        put(e, f"{col}{er}", f"=SUM({col}7:{col}{el})", BLACK, USD, bold=True)
        put(e, f"{col}{er+1}", "=Assumptions!$B$13", GREEN, USD)
        put(e, f"{col}{er+2}", f"=IF({col}{er+1}=0,0,{col}{er}/{col}{er+1})", BLACK, PCT)
    widths(e, {"A": 38, "B": 10, "C": 14, "D": 12, "E": 18, "F": 13, "G": 13, "H": 13, "I": 13, "J": 13})

    # ---------------- Summary
    s = wb.create_sheet("Summary")
    title(s, "IRRBB Summary vs Limits", "Feeds: Annual Report (Market Risk) and Pillar 3 (IRRBB)")
    header(s, 5, ["Scenario", "Shock (bps)", "Projected NII ($mm)", "Change in NII ($mm)", "Change in NII (%)",
                  "NII limit status", "Change in EVE ($mm)", "EVE % of Tier 1", "EVE limit status"])
    for j in range(5):
        r = 6 + j
        ncol = get_column_letter(4 + j)
        ecol = get_column_letter(6 + j)
        put(s, f"A{r}", f"=Assumptions!A{17+j}", GREEN)
        put(s, f"B{r}", f"=Assumptions!B{17+j}", GREEN, BPS)
        put(s, f"C{r}", f"=NII_Projection!{ncol}{nii_rows['proj']}", GREEN, USD)
        put(s, f"D{r}", f"=NII_Projection!{ncol}{nii_rows['delta']}", GREEN, USD)
        put(s, f"E{r}", f"=NII_Projection!{ncol}{nii_rows['pct']}", GREEN, PCT)
        put(s, f"F{r}", f'=IF(E{r}<-Assumptions!$B$11,"BREACH","Within limit")', BLACK)
        put(s, f"G{r}", f"=EVE!{ecol}{er}", GREEN, USD)
        put(s, f"H{r}", f"=EVE!{ecol}{er+2}", GREEN, PCT)
        put(s, f"I{r}", f'=IF(H{r}<-Assumptions!$B$12,"BREACH","Within limit")', BLACK)
    widths(s, {"A": 14, "B": 12, "C": 18, "D": 18, "E": 16, "F": 16, "G": 18, "H": 16, "I": 16})
    wb.move_sheet("Summary", offset=-4)
    path = f"{OUT}/{m['file']}"
    wb.save(path)
    return path


# =============================================================== CECL MODEL
CECL_PARAMS = {
    # segment: life_yrs, base_pd, lgd, pd_elasticity (per 1pp unemployment), lgd_sens, price driver
    "Credit Card": (1.70, 0.0374, 0.880, 0.165, 0.00, "None"),
    "Residential Mortgage": (6.50, 0.0035, 0.140, 0.140, 1.80, "HPI"),
    "Auto": (2.40, 0.0104, 0.420, 0.120, 0.40, "None"),
    "Commercial Real Estate": (3.80, 0.0137, 0.380, 0.150, 1.60, "CRE"),
    "Commercial & Industrial": (2.30, 0.0150, 0.410, 0.135, 0.00, "None"),
    "Other Consumer & Wholesale": (2.00, 0.0149, 0.330, 0.100, 0.00, "None"),
}


def build_cecl(overlays=None):
    m = B.MODELS[1]
    wb = Workbook()
    readme(wb, m,
           [("Scenarios", "Probability-weighted macroeconomic scenarios (Upside / Baseline / Downside)."),
            ("Segment_Inputs", "Exposure at default, remaining life, PD and LGD parameters and macro sensitivities."),
            ("ECL_Calc", "Scenario-conditional lifetime expected credit loss by segment."),
            ("Allowance_Summary", "Probability-weighted allowance, qualitative overlay, reconciliation to the "
                                  "reported allowance and coverage metrics."),
            ("Sensitivity", "Allowance under 100% weight on each single scenario.")],
           ["Scenario PD = base annual PD x (1 + elasticity x (scenario peak unemployment - baseline "
            "unemployment) x 100), floored at 0.",
            "Lifetime PD = 1 - (1 - scenario PD) ^ remaining life (years).",
            "Scenario LGD = base LGD x (1 - LGD sensitivity x collateral price change), capped at 100%. "
            "Mortgages use the house price index (HPI); CRE uses the commercial property price index.",
            "Scenario ECL = EAD x lifetime PD x scenario LGD. Modeled allowance = sum of scenario weight x ECL.",
            "Total allowance = modeled allowance + qualitative overlay approved by the Allowance Committee."],
           ["Loan balances (EAD) and remaining life from API-11 Loan Servicing API.",
            "PD and LGD parameters from API-10 Credit Risk Scoring API (pool-level, refreshed quarterly).",
            "Scenario paths and weights from API-14 Macroeconomic Scenario API (approved by the Scenario "
            "Committee)."],
           ["Allowance for loan losses by portfolio segment and in total.",
            "Allowance coverage ratio (allowance / loans) and coverage of annual net charge-offs.",
            "Sensitivity of the allowance to a 100% Downside weighting."],
           ["Simplified pool-level approach for demonstration; production model uses loan-level cash flows.",
            "Unfunded commitments (off-balance-sheet allowance) are excluded.",
            "Qualitative overlays must be documented and fall within +/-15% of the modeled allowance per segment.",
            "Scenario weights must sum to 100% (check cell on Scenarios sheet)."])

    sc = wb.create_sheet("Scenarios")
    title(sc, "Macroeconomic Scenarios", "Source: API-14 Macroeconomic Scenario API, scenario set MSC-2025Q4")
    header(sc, 5, ["Scenario", "Weight", "Peak unemployment", "Real GDP growth 2026", "House price index chg",
                   "CRE price index chg"])
    for i, (name, vals) in enumerate(B.MACRO_SCENARIOS.items()):
        r = 6 + i
        put(sc, f"A{r}", name, BOLD)
        put(sc, f"B{r}", vals[0], BLUE, PCT, YELLOW)
        for col, v in zip("CDEF", vals[1:]):
            put(sc, f"{col}{r}", v, BLUE, PCT)
    put(sc, "A9", "Weight check", BOLD)
    put(sc, "B9", "=SUM(B6:B8)", BLACK, PCT)
    put(sc, "C9", '=IF(ABS(B9-1)<0.0001,"OK","WEIGHTS MUST SUM TO 100%")', BLACK)
    put(sc, "A11", "Baseline unemployment (anchor)", BOLD)
    put(sc, "B11", "=C7", GREEN, PCT)
    widths(sc, {"A": 30, "B": 12, "C": 18, "D": 18, "E": 18, "F": 18})

    si = wb.create_sheet("Segment_Inputs")
    title(si, "Portfolio Segment Inputs", "Balances $mm at 2025-12-31")
    header(si, 5, ["Segment", "Owning LOB", "EAD / balance ($mm)", "Remaining life (yrs)", "Base annual PD",
                   "Base LGD", "PD elasticity (per 1pp unemp.)", "LGD sensitivity to price", "Price driver",
                   "FY2025 net charge-offs ($mm)", "Reported allowance ($mm)"])
    for i, (seg, bal, allow, lob) in enumerate(B.CECL_SEGMENTS):
        r = 6 + i
        life, pd, lgd, el, ls, drv = CECL_PARAMS[seg]
        put(si, f"A{r}", seg, BOLD)
        put(si, f"B{r}", lob, BLUE)
        put(si, f"C{r}", bal, BLUE, USD)
        put(si, f"D{r}", life, BLUE, NUM2)
        put(si, f"E{r}", pd, BLUE, PCT2, YELLOW)
        put(si, f"F{r}", lgd, BLUE, PCT)
        put(si, f"G{r}", el, BLUE, '0.000')
        put(si, f"H{r}", ls, BLUE, NUM2)
        put(si, f"I{r}", drv, BLUE)
        put(si, f"J{r}", B.NCO_BY_SEGMENT[seg], BLUE, USD)
        put(si, f"K{r}", allow, BLUE, USD,
            note="Reported in Annual Report Note 6 (Allowance for Credit Losses).")
    put(si, "A12", "Total", BOLD)
    for col in "CJK":
        put(si, f"{col}12", f"=SUM({col}6:{col}11)", BLACK, USD, bold=True)
    widths(si, {"A": 28, "B": 11, "C": 16, "D": 12, "E": 12, "F": 10, "G": 14, "H": 14, "I": 10, "J": 16, "K": 16})

    ec = wb.create_sheet("ECL_Calc")
    title(ec, "Scenario-Conditional Lifetime ECL", "One block per scenario; rows = portfolio segments")
    scen_names = list(B.MACRO_SCENARIOS.keys())
    block_rows = {}
    r0 = 5
    for k, sname in enumerate(scen_names):
        srow = 6 + k  # row on Scenarios sheet
        top = r0 + k * 10
        ec.cell(row=top, column=1, value=f"Scenario: {sname}").font = TITLE
        header(ec, top + 1, ["Segment", "Scenario PD (annual)", "Lifetime PD", "Collateral price chg",
                             "Scenario LGD", "EAD ($mm)", "Lifetime ECL ($mm)"])
        for i in range(6):
            r = top + 2 + i
            sr = 6 + i
            put(ec, f"A{r}", f"=Segment_Inputs!A{sr}", GREEN)
            put(ec, f"B{r}",
                f"=MAX(0,Segment_Inputs!E{sr}*(1+Segment_Inputs!G{sr}*(Scenarios!$C${srow}-Scenarios!$B$11)*100))",
                BLACK, PCT2)
            put(ec, f"C{r}", f"=1-(1-B{r})^Segment_Inputs!D{sr}", BLACK, PCT2)
            put(ec, f"D{r}",
                f'=IF(Segment_Inputs!I{sr}="HPI",Scenarios!$E${srow},IF(Segment_Inputs!I{sr}="CRE",Scenarios!$F${srow},0))',
                BLACK, PCT)
            put(ec, f"E{r}", f"=MIN(1,MAX(0,Segment_Inputs!F{sr}*(1-Segment_Inputs!H{sr}*D{r})))", BLACK, PCT)
            put(ec, f"F{r}", f"=Segment_Inputs!C{sr}", GREEN, USD)
            put(ec, f"G{r}", f"=F{r}*C{r}*E{r}", BLACK, USD)
        tot = top + 8
        put(ec, f"A{tot}", f"Total - {sname}", BOLD)
        put(ec, f"G{tot}", f"=SUM(G{top+2}:G{top+7})", BLACK, USD, bold=True)
        block_rows[sname] = top + 2
    widths(ec, {"A": 30, "B": 16, "C": 13, "D": 16, "E": 13, "F": 14, "G": 18})

    al = wb.create_sheet("Allowance_Summary")
    title(al, "Allowance for Loan Losses - Summary and Reconciliation", "Feeds Annual Report Note 6 and the earnings supplement")
    header(al, 5, ["Segment", "Upside ECL", "Baseline ECL", "Downside ECL", "Weighted (modeled) ($mm)",
                   "Qualitative overlay ($mm)", "Total allowance ($mm)", "Reported allowance ($mm)", "Difference",
                   "Overlay % of modeled", "Allowance / loans", "Coverage of NCOs (yrs)"])
    if overlays is None:
        overlays = [0] * 6
    for i in range(6):
        r = 6 + i
        sr = 6 + i
        put(al, f"A{r}", f"=Segment_Inputs!A{sr}", GREEN)
        for col, sname in zip("BCD", scen_names):
            put(al, f"{col}{r}", f"=ECL_Calc!G{block_rows[sname]+i}", GREEN, USD)
        put(al, f"E{r}", f"=B{r}*Scenarios!$B$6+C{r}*Scenarios!$B$7+D{r}*Scenarios!$B$8", BLACK, USD)
        put(al, f"F{r}", overlays[i], BLUE, USD, YELLOW,
            note="Management qualitative adjustment approved by the Allowance Committee (Dec-2025): "
                 "captures model limitations, concentration and emerging-risk factors.")
        put(al, f"G{r}", f"=E{r}+F{r}", BLACK, USD, bold=True)
        put(al, f"H{r}", f"=Segment_Inputs!K{sr}", GREEN, USD)
        put(al, f"I{r}", f"=G{r}-H{r}", BLACK, USD)
        put(al, f"J{r}", f"=IF(E{r}=0,0,F{r}/E{r})", BLACK, PCT)
        put(al, f"K{r}", f"=IF(Segment_Inputs!C{sr}=0,0,G{r}/Segment_Inputs!C{sr})", BLACK, PCT2)
        put(al, f"L{r}", f"=IF(Segment_Inputs!J{sr}=0,0,G{r}/Segment_Inputs!J{sr})", BLACK, NUM2)
    put(al, "A12", "Total", BOLD)
    for col in "BCDEFGHI":
        put(al, f"{col}12", f"=SUM({col}6:{col}11)", BLACK, USD, bold=True)
    put(al, "J12", "=IF(E12=0,0,F12/E12)", BLACK, PCT)
    put(al, "K12", "=IF(Segment_Inputs!C12=0,0,G12/Segment_Inputs!C12)", BLACK, PCT2)
    put(al, "L12", "=IF(Segment_Inputs!J12=0,0,G12/Segment_Inputs!J12)", BLACK, NUM2)
    put(al, "A14", "Overlay control check (|overlay| <= 15% of modeled per segment)", BOLD)
    put(al, "F14", '=IF(SUMPRODUCT(--(ABS(J6:J11)>0.15))=0,"OK","REVIEW")', BLACK)
    widths(al, {"A": 30, "B": 12, "C": 12, "D": 12, "E": 16, "F": 16, "G": 16, "H": 16, "I": 11, "J": 12, "K": 12, "L": 12})

    se = wb.create_sheet("Sensitivity")
    title(se, "Allowance Sensitivity to Scenario Weighting", "Modeled allowance plus the same qualitative overlay")
    header(se, 5, ["Weighting", "Modeled ECL ($mm)", "Plus overlay ($mm)", "Change vs reported ($mm)", "Change (%)"])
    rows = [("100% Upside", "=Allowance_Summary!B12"), ("100% Baseline", "=Allowance_Summary!C12"),
            ("100% Downside", "=Allowance_Summary!D12"), ("Probability-weighted (reported)", "=Allowance_Summary!E12")]
    for i, (lab, f) in enumerate(rows):
        r = 6 + i
        put(se, f"A{r}", lab, BOLD)
        put(se, f"B{r}", f, GREEN, USD)
        put(se, f"C{r}", f"=B{r}+Allowance_Summary!$F$12", BLACK, USD)
        put(se, f"D{r}", f"=C{r}-Allowance_Summary!$H$12", BLACK, USD)
        put(se, f"E{r}", f"=IF(Allowance_Summary!$H$12=0,0,D{r}/Allowance_Summary!$H$12)", BLACK, PCT)
    widths(se, {"A": 32, "B": 18, "C": 18, "D": 22, "E": 12})
    wb.move_sheet("Allowance_Summary", offset=-4)
    path = f"{OUT}/{m['file']}"
    wb.save(path)
    return path


# =============================================================== CAPITAL MODEL
QUARTERS = ["Q3 2026", "Q4 2026", "Q1 2027", "Q2 2027", "Q3 2027", "Q4 2027", "Q1 2028", "Q2 2028", "Q3 2028"]
STRESS = {  # per quarter: PPNR, provisions, trading/counterparty loss, RWA growth
    "ppnr": [7_400, 7_100, 6_800, 6_600, 6_500, 6_600, 6_800, 7_000, 7_200],
    "provision": [7_300, 6_800, 6_100, 5_300, 4_400, 3_700, 3_200, 2_900, 2_700],
    "trading_loss": [5_400, 0, 0, 0, 0, 0, 0, 0, 0],
    "rwa_growth": [0.022, 0.015, 0.008, 0.004, 0.0, -0.002, -0.004, -0.004, -0.004],
}


def build_capital():
    m = B.MODELS[2]
    wb = Workbook()
    readme(wb, m,
           [("Assumptions", "Starting capital position (Q2 2026), distribution plan, growth and requirement stack."),
            ("Baseline_Projection", "Nine-quarter CET1 projection under the firm's baseline plan."),
            ("Stress_Projection", "Nine-quarter projection under the severely adverse scenario."),
            ("Summary", "Minimum ratios, capital headroom and an indicative stress capital buffer (SCB).")],
           ["CET1(t) = CET1(t-1) + net income - preferred dividends - common dividends - share repurchases "
            "+ other CET1 movements.",
            "Common dividends = shares outstanding x dividend per share; shares fall by buyback dollars / "
            "assumed share price.",
            "RWA(t) = RWA(t-1) x (1 + quarterly growth).",
            "Stress: pre-tax income = PPNR - provisions - trading & counterparty losses; tax at the effective "
            "rate (tax benefit recognised on losses); buybacks suspended; dividends held flat.",
            "Indicative SCB = start CET1 ratio - minimum stressed CET1 ratio + four quarters of planned "
            "dividends / starting RWA, floored at 2.5%."],
           ["Starting CET1 capital and RWA from API-16 Regulatory Reporting API (FR Y-9C / FFIEC 101 extract).",
            "Scenario paths from API-14 Macroeconomic Scenario API (severely adverse, CCAR 2026 analogue).",
            "Liquidity constraints on distributions cross-checked against API-15 Treasury Liquidity Positions API.",
            "Distribution plan (dividends, buybacks) approved by the Board Capital Committee."],
           ["Quarter-end CET1 ratio path versus management target and regulatory requirement.",
            "Minimum stressed CET1 ratio and indicative SCB.",
            "Capacity for additional buybacks at the management target."],
           ["Management target CET1 ratio: 13.0%; regulatory requirement = 4.5% + SCB + G-SIB surcharge.",
            "Simplified: no AOCI volatility modelling, no deferred-tax-asset threshold deductions.",
            "Stress losses are top-down inputs supplied by the enterprise stress testing program."])

    a = wb.create_sheet("Assumptions")
    title(a, "Assumptions", "Starting point: Q2 2026 actuals")
    rows = [
        ("Starting CET1 capital ($mm)", B.Q2_CAPITAL["CET1 capital"][0], USD, "API-16, Q2 2026 regulatory capital"),
        ("Starting standardized RWA ($mm)", B.Q2_CAPITAL["Standardized RWA"][0], USD, "API-16, Q2 2026"),
        ("Starting common shares outstanding (mm)", 1_381.0, NUM1, "Transfer agent, 2026-06-30"),
        ("Baseline quarterly net income ($mm)", 6_600, USD, "Financial plan 2026-2028"),
        ("Net income growth per quarter", 0.008, PCT, "Financial plan"),
        ("Preferred dividends per quarter ($mm)", 260, USD, "Contractual"),
        ("Common dividend per share per quarter ($)", 1.15, '$0.00', "Board-declared; flat in projection"),
        ("Share repurchases per quarter ($mm)", 3_000, USD, "Board-authorised program ($30bn, 2025-2028)"),
        ("Assumed share price ($)", 262.0, '$0.00', "Flat assumption for share count roll-forward"),
        ("Baseline RWA growth per quarter", 0.010, PCT, "Balance sheet plan"),
        ("Other CET1 movements per quarter ($mm)", -150, USD, "AOCI, deductions, employee stock plans (net)"),
        ("Effective tax rate", 0.22, PCT, "FY2025 effective rate"),
        ("CET1 regulatory minimum", B.CAP_REQ["minimum"], PCT, "12 CFR 217"),
        ("Stress capital buffer (current)", B.CAP_REQ["stress_capital_buffer"], PCT, "Effective 2025-10-01"),
        ("G-SIB surcharge", B.CAP_REQ["gsib_surcharge"], PCT, "Method 2"),
        ("Management target CET1 ratio", B.CAP_REQ["management_target"], PCT, "Board Capital Committee"),
    ]
    header(a, 5, ["Assumption", "Value", "Source / note"])
    for i, (k, v, fmt, note) in enumerate(rows):
        r = 6 + i
        put(a, f"A{r}", k, BOLD)
        put(a, f"B{r}", v, BLUE, fmt, YELLOW if i in (3, 7, 9, 15) else None)
        put(a, f"C{r}", note, SUB)
    put(a, "A23", "Regulatory CET1 requirement", BOLD)
    put(a, "B23", "=B18+B19+B20", BLACK, PCT, bold=True)
    widths(a, {"A": 44, "B": 14, "C": 52})
    # assumption cells
    A_ = {k: f"Assumptions!$B${6+i}" for i, k in enumerate(
        ["cet1", "rwa", "shares", "ni", "nig", "pref", "dps", "bb", "px", "rwag", "other", "tax",
         "min", "scb", "gsib", "target"])}
    A_["req"] = "Assumptions!$B$23"

    def projection(ws, stress):
        title(ws, "Severely Adverse Projection" if stress else "Baseline Projection",
              "Quarterly, $mm" + (" - stress losses from enterprise stress testing program" if stress else ""))
        header(ws, 5, ["Line item", "Q2 2026 (actual)"] + QUARTERS)
        labels = ["Beginning CET1 capital"]
        if stress:
            labels += ["Pre-provision net revenue (PPNR)", "Provision for credit losses",
                       "Trading & counterparty losses", "Pre-tax income", "Income tax (benefit)"]
        labels += ["Net income", "Preferred dividends", "Common dividends", "Share repurchases",
                   "Other CET1 movements", "Ending CET1 capital", "RWA growth", "Standardized RWA",
                   "CET1 ratio", "Shares outstanding (mm)", "Headroom vs management target ($mm)",
                   "Headroom vs regulatory requirement ($mm)"]
        R = {lab: 6 + i for i, lab in enumerate(labels)}
        for lab, r in R.items():
            put(ws, f"A{r}", lab, BOLD)
        # actual column B
        put(ws, f"B{R['Ending CET1 capital']}", f"={A_['cet1']}", GREEN, USD)
        put(ws, f"B{R['Standardized RWA']}", f"={A_['rwa']}", GREEN, USD)
        put(ws, f"B{R['CET1 ratio']}", f"=B{R['Ending CET1 capital']}/B{R['Standardized RWA']}", BLACK, PCT2)
        put(ws, f"B{R['Shares outstanding (mm)']}", f"={A_['shares']}", GREEN, NUM1)
        for j, q in enumerate(QUARTERS):
            c = get_column_letter(3 + j)
            p = get_column_letter(2 + j)
            put(ws, f"{c}{R['Beginning CET1 capital']}", f"={p}{R['Ending CET1 capital']}", BLACK, USD)
            if stress:
                put(ws, f"{c}{R['Pre-provision net revenue (PPNR)']}", STRESS["ppnr"][j], BLUE, USD)
                put(ws, f"{c}{R['Provision for credit losses']}", -STRESS["provision"][j], BLUE, USD)
                put(ws, f"{c}{R['Trading & counterparty losses']}", -STRESS["trading_loss"][j], BLUE, USD)
                put(ws, f"{c}{R['Pre-tax income']}",
                    f"=SUM({c}{R['Pre-provision net revenue (PPNR)']}:{c}{R['Trading & counterparty losses']})", BLACK, USD)
                put(ws, f"{c}{R['Income tax (benefit)']}", f"=-{c}{R['Pre-tax income']}*{A_['tax']}", BLACK, USD)
                put(ws, f"{c}{R['Net income']}", f"={c}{R['Pre-tax income']}+{c}{R['Income tax (benefit)']}", BLACK, USD)
                put(ws, f"{c}{R['Share repurchases']}", 0, BLUE, USD, note="Buybacks suspended under stress (CCAR convention).")
                put(ws, f"{c}{R['RWA growth']}", STRESS["rwa_growth"][j], BLUE, PCT)
            else:
                put(ws, f"{c}{R['Net income']}", f"={A_['ni']}*(1+{A_['nig']})^{j}", BLACK, USD)
                put(ws, f"{c}{R['Share repurchases']}", f"=-{A_['bb']}", GREEN, USD)
                put(ws, f"{c}{R['RWA growth']}", f"={A_['rwag']}", GREEN, PCT)
            put(ws, f"{c}{R['Preferred dividends']}", f"=-{A_['pref']}", GREEN, USD)
            put(ws, f"{c}{R['Common dividends']}", f"=-{p}{R['Shares outstanding (mm)']}*{A_['dps']}", BLACK, USD)
            put(ws, f"{c}{R['Other CET1 movements']}", f"={A_['other']}", GREEN, USD)
            put(ws, f"{c}{R['Ending CET1 capital']}",
                f"={c}{R['Beginning CET1 capital']}+{c}{R['Net income']}+{c}{R['Preferred dividends']}"
                f"+{c}{R['Common dividends']}+{c}{R['Share repurchases']}+{c}{R['Other CET1 movements']}",
                BLACK, USD, bold=True)
            put(ws, f"{c}{R['Standardized RWA']}", f"={p}{R['Standardized RWA']}*(1+{c}{R['RWA growth']})", BLACK, USD)
            put(ws, f"{c}{R['CET1 ratio']}", f"=IF({c}{R['Standardized RWA']}=0,0,{c}{R['Ending CET1 capital']}/{c}{R['Standardized RWA']})",
                BLACK, PCT2, bold=True)
            put(ws, f"{c}{R['Shares outstanding (mm)']}",
                f"={p}{R['Shares outstanding (mm)']}+{c}{R['Share repurchases']}/{A_['px']}", BLACK, NUM1)
            put(ws, f"{c}{R['Headroom vs management target ($mm)']}",
                f"={c}{R['Ending CET1 capital']}-{A_['target']}*{c}{R['Standardized RWA']}", BLACK, USD)
            put(ws, f"{c}{R['Headroom vs regulatory requirement ($mm)']}",
                f"={c}{R['Ending CET1 capital']}-{A_['req']}*{c}{R['Standardized RWA']}", BLACK, USD)
        widths(ws, {"A": 40, **{get_column_letter(k): 13 for k in range(2, 12)}})
        ws.freeze_panes = "B6"
        return R

    bp = wb.create_sheet("Baseline_Projection")
    RB = projection(bp, False)
    sp = wb.create_sheet("Stress_Projection")
    RS = projection(sp, True)

    s = wb.create_sheet("Summary")
    title(s, "Capital Planning Summary", "Feeds Annual Report (Capital Risk Management) and Pillar 3 (stress testing)")
    header(s, 5, ["Metric", "Value", "Note"])
    lastc = get_column_letter(2 + len(QUARTERS))
    rows = [
        ("Starting CET1 ratio (Q2 2026)", f"=Baseline_Projection!B{RB['CET1 ratio']}", PCT2, ""),
        ("Ending CET1 ratio - baseline (Q3 2028)", f"=Baseline_Projection!{lastc}{RB['CET1 ratio']}", PCT2, ""),
        ("Minimum CET1 ratio - baseline", f"=MIN(Baseline_Projection!C{RB['CET1 ratio']}:{lastc}{RB['CET1 ratio']})", PCT2, ""),
        ("Minimum CET1 ratio - severely adverse", f"=MIN(Stress_Projection!C{RS['CET1 ratio']}:{lastc}{RS['CET1 ratio']})", PCT2, ""),
        ("Peak-to-trough CET1 decline (stress)", "=B6-B9", PCT2, "Starting ratio minus minimum stressed ratio"),
        ("Four quarters of planned dividends / starting RWA",
         f"=-SUM(Baseline_Projection!C{RB['Common dividends']}:F{RB['Common dividends']})/Baseline_Projection!B{RB['Standardized RWA']}",
         PCT2, "Dividend add-on in SCB formula"),
        ("Indicative stress capital buffer (SCB)", "=MAX(0.025,B10+B11)", PCT2, "Floored at 2.5%"),
        ("Regulatory CET1 requirement (current SCB)", f"={A_['req']}", PCT2, "4.5% + SCB + G-SIB surcharge"),
        ("Regulatory CET1 requirement (indicative SCB)", f"={A_['min']}+B12+{A_['gsib']}", PCT2, ""),
        ("Management target", f"={A_['target']}", PCT2, ""),
        ("Cumulative buybacks over horizon ($mm)", f"=-SUM(Baseline_Projection!C{RB['Share repurchases']}:{lastc}{RB['Share repurchases']})", USD, ""),
        ("Excess CET1 vs target at Q3 2028 ($mm)", f"=Baseline_Projection!{lastc}{RB['Headroom vs management target ($mm)']}", USD,
         "Additional distribution capacity at the target"),
        ("Stress minimum above requirement?", '=IF(B9>=B13,"YES","NO - capital action required")', None, ""),
    ]
    for i, (k, f, fmt, note) in enumerate(rows):
        r = 6 + i
        put(s, f"A{r}", k, BOLD)
        put(s, f"B{r}", f, GREEN if "!" in f else BLACK, fmt, bold=True)
        put(s, f"C{r}", note, SUB)
    widths(s, {"A": 50, "B": 16, "C": 44})
    wb.move_sheet("Summary", offset=-3)
    path = f"{OUT}/{m['file']}"
    wb.save(path)
    return path


if __name__ == "__main__":
    import json, os
    os.makedirs(OUT, exist_ok=True)
    ov = None
    if os.path.exists("cecl_overlays.json"):
        ov = json.load(open("cecl_overlays.json"))
    print(build_nii())
    print(build_cecl(ov))
    print(build_capital())
